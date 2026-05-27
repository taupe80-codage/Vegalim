"""
patch_v32_remove_synthetic_ids.py
==================================
Supprime de ingredients_v32.json tous les variants dont le source_id
est un ID synthétique généré par build_cnf_full_v10.py (CNF >= 500000)
ou tout autre ID non présent dans les sources brutes officielles.

Ces IDs n'existent pas dans le CNF réel — ils avaient été créés par
l'ancien pipeline flat pour représenter des combinaisons cuites/salées.
"""
import csv, json, sys
from pathlib import Path
from collections import defaultdict

sys.stdout.reconfigure(encoding='utf-8')

ROOT    = Path(__file__).parents[2]
DATA    = ROOT / 'backend/data'
V32_PATH = DATA / 'ingredients/ingredients_v32.json'
RAW_DIR  = DATA / 'nutrition/raw'

# ── IDs officiels valides ──────────────────────────────────────────────────────
print('Chargement IDs officiels...')

import openpyxl
wb = openpyxl.load_workbook(RAW_DIR / 'Table_Ciqual_2025_FR_2025_11_03.xlsx', read_only=True)
ws = wb.active
next(ws.iter_rows(max_row=1))
ciq_ids = set()
for row in ws.iter_rows(min_row=2, values_only=True):
    c = row[6]
    if c is not None:
        try: ciq_ids.add(int(c))
        except: pass
wb.close()
print(f'  CIQUAL : {len(ciq_ids)} IDs valides')

usda_data = json.loads(
    (RAW_DIR / 'FoodData_Central_foundation_food_json_2025-12-18.json').read_text(encoding='utf-8')
)
usda_ids = {int(f['fdcId']) for f in usda_data.get('FoundationFoods', []) if f.get('fdcId')}
print(f'  USDA   : {len(usda_ids)} IDs valides')

cnf_ids = set()
with open(RAW_DIR / 'cnf/FOOD_NAME.csv', encoding='utf-8-sig', errors='replace') as f:
    for row in csv.DictReader(f):
        fc = row.get('FoodCode') or row.get('FoodID') or ''
        try: cnf_ids.add(int(fc))
        except: pass
print(f'  CNF    : {len(cnf_ids)} IDs valides')

def is_valid(src: str, sid: int) -> bool:
    if src == 'CIQUAL': return sid in ciq_ids
    if src == 'USDA':   return sid in usda_ids
    if src == 'CNF':    return sid in cnf_ids
    return False

# ── Patch v32 ─────────────────────────────────────────────────────────────────
print('\nApplication du patch...')
v32 = json.loads(V32_PATH.read_text(encoding='utf-8'))

removed      = []           # liste des variants supprimés
ig_emptied   = []           # groupes qui n'ont plus aucun variant après suppression
stats = {'removed': 0, 'kept': 0, 'ig_empty': 0}

# Grouper par catégorie pour le rapport
by_cat = defaultdict(list)

for cat in v32.get('categories', []):
    cat1 = cat.get('label', '')
    for sub in cat.get('subcategories', []):
        cat2 = sub.get('label', '')
        for ig in sub.get('ingredient_groups', []):
            ig_id   = ig.get('id', '')
            en      = ig.get('canonical_name_en', '')
            kept    = []
            dropped = []

            for vr in ig.get('variants', []):
                src = (vr.get('source') or '').upper()
                sid_raw = vr.get('source_id')
                if sid_raw is None:
                    kept.append(vr)
                    continue
                sid = int(sid_raw)
                if is_valid(src, sid):
                    kept.append(vr)
                    stats['kept'] += 1
                else:
                    dropped.append({
                        'vr_id': vr.get('id',''),
                        'src': src, 'source_id': sid,
                        'ig_id': ig_id, 'en': en[:40],
                        'cat2': cat2,
                    })
                    stats['removed'] += 1
                    by_cat[cat2].append(f"{src}#{sid} ({vr.get('id','')})")

            ig['variants'] = kept

            if dropped and not kept:
                ig_emptied.append({'ig_id': ig_id, 'en': en, 'cat2': cat2})
                stats['ig_empty'] += 1

# ── Rapport ────────────────────────────────────────────────────────────────────
print()
print('=' * 60)
print('SUPPRESSION DES IDs SYNTHÉTIQUES — RÉSULTAT')
print('=' * 60)
print(f'  Variants supprimés    : {stats["removed"]}')
print(f'  Variants conservés    : {stats["kept"]}')
print(f'  Groupes vidés (0 var) : {stats["ig_empty"]}')
print()

print('── Par catégorie (top) ──────────────────────────────────────')
for cat2, items in sorted(by_cat.items(), key=lambda x: -len(x[1]))[:20]:
    print(f'  {cat2:<30} : {len(items)} supprimé(s)')

if ig_emptied:
    print()
    print('── Groupes maintenant sans variant ──────────────────────────')
    for ig in ig_emptied:
        print(f"  {ig['ig_id']} | {ig['cat2']:<25} | {ig['en']}")

# ── Sauvegarde ─────────────────────────────────────────────────────────────────
V32_PATH.write_text(json.dumps(v32, ensure_ascii=False, indent=2), encoding='utf-8')
print(f'\n  ✅ ingredients_v32.json mis à jour')
