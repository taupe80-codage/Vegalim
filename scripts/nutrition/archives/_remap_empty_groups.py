"""
Cherche les vrais FoodCode CNF et alim_code CIQUAL pour les groupes v32 vidés.
Stratégie : matching par nom EN/FR dans les sources brutes.
"""
import csv, json, re, sys
from pathlib import Path
from collections import defaultdict

sys.stdout.reconfigure(encoding='utf-8')

ROOT     = Path(__file__).parents[2]
DATA     = ROOT / 'backend/data'
V32_PATH = DATA / 'ingredients/ingredients_v32.json'
RAW_DIR  = DATA / 'nutrition/raw'

def slug(s):
    s = (s or '').lower().strip()
    for c, r in [('é','e'),('è','e'),('ê','e'),('à','a'),('â','a'),
                  ('ù','u'),('û','u'),('î','i'),('ï','i'),('ô','o'),('ç','c')]:
        s = s.replace(c, r)
    s = re.sub(r"['\u2019]", '', s)
    s = re.sub(r'[^a-z0-9]+', ' ', s)
    return s.strip()

def token_sim(a, b):
    ta = set(slug(a).split()) - {'de','du','des','le','la','les','et','au','aux','or','and','the','with'}
    tb = set(slug(b).split()) - {'de','du','des','le','la','les','et','au','aux','or','and','the','with'}
    if not ta or not tb: return 0.0
    return len(ta & tb) / max(len(ta), len(tb))

# Charger les sources brutes (noms uniquement)
import openpyxl
print('Chargement sources...')
wb = openpyxl.load_workbook(RAW_DIR / 'Table_Ciqual_2025_FR_2025_11_03.xlsx', read_only=True)
ws = wb.active
next(ws.iter_rows(max_row=1))
ciq_names = {}  # alim_code → (name_fr, name_grp)
for row in ws.iter_rows(min_row=2, values_only=True):
    c = row[6]
    if c is not None:
        try: ciq_names[int(c)] = (str(row[7] or ''), str(row[3] or ''))
        except: pass
wb.close()
print(f'  CIQUAL: {len(ciq_names)}')

cnf_names = {}  # FoodCode → (name_en, name_fr)
with open(RAW_DIR / 'cnf/FOOD_NAME.csv', encoding='utf-8-sig', errors='replace') as f:
    for row in csv.DictReader(f):
        fc = row.get('FoodCode') or row.get('FoodID') or ''
        try:
            cnf_names[int(fc)] = (row.get('FoodDescription',''), row.get('FoodDescriptionF',''))
        except: pass
print(f'  CNF: {len(cnf_names)}')

# Groupes vidés (variant_count == 0)
v32 = json.loads(V32_PATH.read_text(encoding='utf-8'))
empty_groups = []
for cat in v32.get('categories', []):
    for sub in cat.get('subcategories', []):
        for ig in sub.get('ingredient_groups', []):
            if not ig.get('variants'):
                empty_groups.append({
                    'ig_id': ig['id'],
                    'en': ig.get('canonical_name_en',''),
                    'fr': ig.get('canonical_name_fr',''),
                    'cat2': sub.get('label',''),
                })

print(f'\nGroupes vidés : {len(empty_groups)}')
print()

# Pour chaque groupe vide, chercher la meilleure correspondance
results = []
for grp in empty_groups:
    en = grp['en']
    fr = grp['fr']
    best_ciq = None
    best_ciq_sim = 0
    best_cnf = None
    best_cnf_sim = 0

    # CIQUAL
    for code, (name_fr, grp_fr) in ciq_names.items():
        s_fr = token_sim(fr, name_fr) if fr else 0
        s_en = token_sim(en, name_fr)
        s = max(s_fr, s_en)
        if s > best_ciq_sim:
            best_ciq_sim = s
            best_ciq = (code, name_fr)

    # CNF
    for fc, (name_en, name_fr_cnf) in cnf_names.items():
        s_en = token_sim(en, name_en)
        s_fr = token_sim(fr, name_fr_cnf) if fr else 0
        s = max(s_en, s_fr)
        if s > best_cnf_sim:
            best_cnf_sim = s
            best_cnf = (fc, name_en)

    results.append({
        **grp,
        'ciq_id': best_ciq[0] if best_ciq_sim >= 0.4 else None,
        'ciq_name': best_ciq[1][:45] if best_ciq_sim >= 0.4 else '—',
        'ciq_sim': round(best_ciq_sim, 2),
        'cnf_id': best_cnf[0] if best_cnf_sim >= 0.4 else None,
        'cnf_name': best_cnf[1][:45] if best_cnf_sim >= 0.4 else '—',
        'cnf_sim': round(best_cnf_sim, 2),
    })

# Afficher les résultats
RECOVERABLE  = [r for r in results if r['ciq_id'] or r['cnf_id']]
UNRECOVERABLE = [r for r in results if not r['ciq_id'] and not r['cnf_id']]

print(f'  Récupérables (correspondance >= 0.4) : {len(RECOVERABLE)}')
print(f'  À supprimer (aucune correspondance)  : {len(UNRECOVERABLE)}')
print()

print('── RÉCUPÉRABLES ──────────────────────────────────────────────────────')
for r in RECOVERABLE:
    src_info = ''
    if r['ciq_id']:
        src_info += f"  CIQUAL#{r['ciq_id']} sim={r['ciq_sim']} \"{r['ciq_name']}\""
    if r['cnf_id']:
        src_info += f"  CNF#{r['cnf_id']} sim={r['cnf_sim']} \"{r['cnf_name']}\""
    print(f"  {r['ig_id']} | {r['cat2']:<22} | {r['en'][:32]:<32} |{src_info}")

print()
print('── À SUPPRIMER DE v32 ────────────────────────────────────────────────')
for r in UNRECOVERABLE:
    print(f"  {r['ig_id']} | {r['cat2']:<22} | {r['en'][:45]}")

# Export JSON
out_path = DATA / 'nutrition/logs/v32_empty_groups_remap.json'
out_path.parent.mkdir(parents=True, exist_ok=True)
out_path.write_text(json.dumps({
    'recoverable': RECOVERABLE,
    'unrecoverable': UNRECOVERABLE,
}, ensure_ascii=False, indent=2), encoding='utf-8')
print(f'\n  Rapport → {out_path}')
