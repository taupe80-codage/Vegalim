"""
patch_ciqual_axes_v32.py
========================
Audit + patch des axes manquants pour les variants CIQUAL dans v32.
Noms CIQUAL en francais — patterns FR avec filtrage faux positifs.

Mode : DRY-RUN par defaut. --apply pour modifier v32.
"""
import json, re, sys, copy, shutil, csv
from datetime import datetime
from pathlib import Path
from collections import defaultdict

sys.stdout.reconfigure(encoding='utf-8')
DRY_RUN = '--apply' not in sys.argv

ROOT  = Path(__file__).parents[2]
DATA  = ROOT / 'backend/data'
V32_P = DATA / 'ingredients/ingredients_v32.json'
XLSX  = DATA / 'nutrition/raw/Table_Ciqual_2025_FR_2025_11_03.xlsx'

# ── Charger CIQUAL ────────────────────────────────────────────────────────────
print('Chargement CIQUAL xlsx...')
try:
    import openpyxl
except ImportError:
    print('[ERREUR] pip install openpyxl'); sys.exit(1)

wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
ws = wb.active
headers = [str(c.value).strip() if c.value else '' for c in next(ws.iter_rows())]

def find_col(candidates):
    for c in candidates:
        if c in headers:
            return headers.index(c)
    return None

idx_code = find_col(['alim_code', 'Code'])
idx_nom  = find_col(['alim_nom_fr', 'alim_nom_fr (CIQUAL)', 'Aliment'])
print(f'  code col={idx_code}  nom col={idx_nom}')

ciqual_names: dict[str, str] = {}
for row in ws.iter_rows(min_row=2, values_only=True):
    code = str(row[idx_code]).strip() if idx_code is not None and row[idx_code] is not None else ''
    nom  = str(row[idx_nom]).strip()  if idx_nom  is not None and row[idx_nom]  is not None else ''
    if code and nom and code != 'None':
        ciqual_names[code] = nom
wb.close()
print(f'  {len(ciqual_names)} entrees CIQUAL')
for k, v in list(ciqual_names.items())[:2]:
    print(f'  ex: {k} -> {v}')

# ── Charger v32 ──────────────────────────────────────────────────────────────
v32 = json.loads(V32_P.read_text(encoding='utf-8'))


# ══════════════════════════════════════════════════════════════════════════════
# FAUX POSITIFS
# ══════════════════════════════════════════════════════════════════════════════
SKIP_RAW_CATS   = {'spices_and_herbs', 'leavening_agents_and_additives', 'fats_and_oils'}
SKIP_DRIED_CATS = {'spices_and_herbs', 'leavening_agents_and_additives'}
SKIP_FRESH_CATS = {'spices_and_herbs'}

# (pattern_nom, axe, raison)
FALSE_POSITIVES = [
    # "lait cru" / "fromage au lait cru" => pasteurisation, pas cuisson
    (r'lait cru|au lait cru|p.te crue',         'etat_cuisson',   'lait cru = pasteurisation, pas etat cuisson'),
    # "fromage frais" => categorie, pas etat thermique
    (r'fromage frais|faisselle',                 'etat_thermique', 'fromage frais = categorie fromagere'),
    # "lait entier" => teneur_MG, pas forme
    (r'lait entier|lait demi',                   'forme',          'entier pour lait = teneur_MG'),
    # grain entier / farine complete => degre raffinage, pas forme
    (r'grain entier|farine.*(compl|int.gral|semi-compl|T\d)|pain.*(compl|int.gral)',
                                                 'forme',          'entier/complet = degre raffinage'),
    # herbe hachee => presentation standard
    (r'(persil|ciboulette|basilic|menthe|coriandre).+hach',
                                                 'forme',          'herbe hachee = presentation standard'),
    # epice en poudre => identite ingredient
    (r'(ail|oignon|moutarde|paprika|piment|curcuma|cannelle|cumin|poivre|curry|gingembre|vanille).+poudre',
                                                 'forme',          'epice en poudre = identite ingredient'),
    # produit allege generique
    (r'produit allege|boisson alleg',            'teneur_MG',      'alleg generique sans valeur MG precise'),
    # nature != cru
    (r'\bnature\b',                              'etat_cuisson',   'nature = sans ajout, pas cru'),
    # freshwater / eau douce
    (r'eau douce|eau de mer',                    'etat_thermique', 'eau douce/mer = habitat, pas etat'),
    # aromatise => pas un etat de cuisson
    (r'aromatis',                                'etat_cuisson',   'aromatise = traitement, pas cuisson'),
    # "sec" dans noms de legumineuses type "haricot sec" = variete, pas etat thermique
    # On laisse passer car c'est bien un etat_thermique discriminant pour les legumineuses
]

def is_fp(nom: str, axe: str, cat_label: str):
    nom_l = nom.lower()
    if axe == 'etat_cuisson' and cat_label in SKIP_RAW_CATS:
        return f'cat={cat_label}: cru/cuit non pertinent'
    if axe == 'etat_thermique' and cat_label in SKIP_DRIED_CATS:
        return f'cat={cat_label}: seche = etat par defaut'
    if axe == 'etat_thermique' and 'frais' in nom_l and cat_label in SKIP_FRESH_CATS:
        return f'cat={cat_label}: frais = etat par defaut herbes'
    for pattern, fp_axe, reason in FALSE_POSITIVES:
        if fp_axe == axe and re.search(pattern, nom_l, re.IGNORECASE):
            return reason
    return None


# ══════════════════════════════════════════════════════════════════════════════
# REGLES DE DETECTION (noms CIQUAL francais)
# ══════════════════════════════════════════════════════════════════════════════
# Note: utilisation de \xe9 (e accent aigu) etc. pour eviter pbm encodage
# Mais Python 3 avec source utf-8 gere bien les accents en raw string

RULES_FR = [
    # ── Etat de cuisson ──────────────────────────────────────────────────────
    (r'\bcru\b|\bnon cuit\b',               'etat_cuisson',    'cru',         'high'),
    (r'\bcuit\b',                           'etat_cuisson',    'cuit',        'high'),
    (r'\bbouilli\b|\bcuit .* eau\b',        'etat_cuisson',    'cuit',        'high'),
    (r'\bcuit vapeur\b|\bvapeur\b',         'etat_cuisson',    'cuit_vapeur', 'high'),
    (r'\bfrit\b',                           'etat_cuisson',    'frit',        'high'),
    (r'\bgrill[e\xE9]\b|\br[o\xF4]ti\b',   'etat_cuisson',    'roti',        'high'),
    (r'\bcuit au four\b',                   'etat_cuisson',    'cuit au four','high'),
    (r'\bpoch[e\xE9]\b',                    'etat_cuisson',    'cuit',        'high'),
    (r'\bbrais[e\xE9]\b',                   'etat_cuisson',    'braise',      'high'),

    # ── Etat thermique ───────────────────────────────────────────────────────
    (r'\bs[e\xE9]ch[e\xE9]\b|\bd[e\xE9]shydrat[e\xE9]\b|\blyophilis[e\xE9]\b',
                                            'etat_thermique',  'sec',         'high'),
    (r'\bsurgel[e\xE9]\b|\bcongel[e\xE9]\b','etat_thermique', 'surgele',     'high'),
    (r'\bfrais\b|\bfra[i\xEE]che\b',        'etat_thermique',  'frais',       'high'),

    # ── Conditionnement ──────────────────────────────────────────────────────
    (r'\ben conserve\b|\bappertis[e\xE9]\b','conditionnement','conserve',    'high'),
    (r'\bau sirop\b|\b. l.huile\b|\bau naturel\b',
                                            'conditionnement','conserve',    'high'),
    (r'\bUHT\b|\blongue conservation\b',    'conditionnement','UHT',         'high'),

    # ── Egouttage ────────────────────────────────────────────────────────────
    (r'\b[e\xE9]goutt[e\xE9]\b',           'egouttage_milieu','egoutte',    'high'),

    # ── Assaisonnement ───────────────────────────────────────────────────────
    (r'\bsal[e\xE9]\b|\bavec sel\b',        'assaisonnement',  'sale',        'high'),
    (r'\bsucr[e\xE9]\b|\bavec sucre\b',     'assaisonnement',  'sucre',       'high'),
    (r'\bsans sel\b',                       'assaisonnement',  'sans_sel',    'high'),
    (r'\bsans sucre\b|\bnon sucr[e\xE9]\b', 'assaisonnement',  'sans_sucre',  'high'),

    # ── Forme ────────────────────────────────────────────────────────────────
    (r'\ben poudre\b|\bpoudre\b',           'forme',           'poudre',      'high'),
    (r'\ben flocons\b|\bflocons\b',         'forme',           'flocons',     'high'),
    (r'\bhach[e\xE9][e\xE9]?\b',            'forme',           'hache',       'medium'),
    (r'\bmoul[ue]\b',                       'forme',           'moulu',       'high'),
    (r'\bpur[e\xE9]e\b|\b[e\xE9]cras[e\xE9]\b','forme',       'puree',       'high'),
    (r'\btranch[e\xE9]\b|\ben tranches\b',  'forme',           'tranche',     'medium'),
    (r'\brap[e\xE9]\b',                     'forme',           'rape',        'high'),
    (r'\bconfit[e\xE9]?\b',                 'forme',           'confit',      'medium'),
    (r'\bconc[e\xE9]ntr[e\xE9]\b',          'forme',           'concentre',   'high'),

    # ── Partie ───────────────────────────────────────────────────────────────
    (r'\bavec peau\b|\bnon [e\xE9]pluche\b|\bnon pel[e\xE9]\b',
                                            'partie',          'avec_peau',   'high'),
    (r'\bsans peau\b|\b[e\xE9]pluche\b|\bpel[e\xE9]\b',
                                            'partie',          'sans_peau',   'high'),
    (r'\bfilet\b',                          'partie',          'filet',       'medium'),
    (r'\bentier\b.*(oeuf|huître|clam)',      'partie',          'entier',      'medium'),

    # ── Teneur MG ────────────────────────────────────────────────────────────
    (r'\b[e\xE9]cr[e\xE9]m[e\xE9]\b',      'teneur_MG',       'ecreme',      'high'),
    (r'\bdemi.[e\xE9]cr[e\xE9]m[e\xE9]\b', 'teneur_MG',       'demi-ecreme', 'high'),
    (r'\ballege\b|\ball[e\xE9]g[e\xE9]\b',  'teneur_MG',       'allege',      'high'),
    (r'\bentier\b.*(lait|yogourt|yaourt|creme)',
                                            'teneur_MG',       'entier',      'high'),
    (r'(lait|creme|yaourt).+\bentier\b',    'teneur_MG',       'entier',      'high'),

    # ── Traitement ───────────────────────────────────────────────────────────
    (r'\bfum[e\xE9][e\xE9]?\b',            'traitement',      'fume',        'medium'),
    (r'\bferment[e\xE9]\b',                 'traitement',      'fermente',    'medium'),
    (r'\btorr[e\xE9]fi[e\xE9]\b',           'traitement',      'torrefie',    'high'),
]


def detect_axes_fr(nom: str, existing: dict, cat_label: str) -> dict:
    suggestions = {}
    nom_l = nom.lower()
    for pattern, axe, valeur, conf in RULES_FR:
        if axe in existing or axe in suggestions:
            continue
        if not re.search(pattern, nom_l, re.IGNORECASE):
            continue
        fp = is_fp(nom, axe, cat_label)
        if fp:
            continue
        suggestions[axe] = (valeur, conf)
    return suggestions


# ── Audit + patch ─────────────────────────────────────────────────────────────
print('\nAnalyse variants CIQUAL...')
stats   = defaultdict(int)
patches: list[dict] = []
errors:  list[dict] = []

v32_patched = copy.deepcopy(v32) if not DRY_RUN else v32

for cat in v32_patched.get('categories', []):
    cat_label = cat.get('label', '')
    for sub in cat.get('subcategories', []):
        for ig in sub.get('ingredient_groups', []):
            ig_id   = ig.get('id', '')
            name_en = ig.get('canonical_name_en', '')
            for vr in ig.get('variants', []):
                if vr.get('source', '').upper() != 'CIQUAL':
                    continue
                stats['total'] += 1
                sid = str(vr.get('source_id', '')).strip()
                nom = ciqual_names.get(sid, '')
                if not nom:
                    stats['not_found'] += 1
                    errors.append({'ig_id': ig_id, 'source_id': sid, 'name_en': name_en})
                    continue
                existing = vr.get('axes', {}) or {}
                sugg = detect_axes_fr(nom, existing, cat_label)
                if not sugg:
                    continue
                stats['patched'] += 1
                stats['axes'] += len(sugg)
                patches.append({
                    'ig_id': ig_id, 'name_en': name_en, 'source_id': sid,
                    'ciqual_nom': nom,
                    'added': {k: v[0] for k, v in sugg.items()},
                    'conf':  {k: v[1] for k, v in sugg.items()},
                    'axes_before': dict(existing),
                })
                if not DRY_RUN:
                    if not isinstance(vr.get('axes'), dict):
                        vr['axes'] = {}
                    for axe, (val, _) in sugg.items():
                        vr['axes'][axe] = val

# ── Ecriture ──────────────────────────────────────────────────────────────────
if not DRY_RUN:
    bk = V32_P.with_name(
        f'ingredients_v32_backup_{datetime.now().strftime("%Y%m%d_%H%M%S")}.json')
    shutil.copy2(V32_P, bk)
    V32_P.write_text(
        json.dumps(v32_patched, ensure_ascii=False, indent=2), encoding='utf-8')
    print(f'[OK] Backup -> {bk.name}')
    print(f'[OK] v32 mis a jour -> {V32_P}')

# ── Rapport ───────────────────────────────────────────────────────────────────
mode = 'DRY-RUN' if DRY_RUN else 'APPLIQUE'
print()
print('=' * 65)
print(f'  PATCH CIQUAL AXES v32 -- {mode}')
print('=' * 65)
print(f'  Variants CIQUAL total     : {stats["total"]}')
print(f'  IDs non trouves           : {stats["not_found"]}')
print(f'  Variants patches          : {stats["patched"]}')
print(f'  Axes ajoutes total        : {stats["axes"]}')
print()

by_axe = defaultdict(int)
for p in patches:
    for axe in p['added']:
        by_axe[axe] += 1
print('  Par axe :')
for axe, n in sorted(by_axe.items(), key=lambda x: -x[1]):
    print(f'    {axe:<25} : {n:>4}')

print()
print('  Apercu (30 premiers) :')
print('  ' + '-' * 62)
for p in patches[:30]:
    print(f'  [{p["ig_id"]}] {p["name_en"][:35]:<35}')
    print(f'    CIQUAL: {p["ciqual_nom"][:65]}')
    for axe, val in p['added'].items():
        cf = p['conf'].get(axe, '')
        print(f'    + {axe}: {val}  [{cf}]')
if len(patches) > 30:
    print(f'  ... et {len(patches) - 30} autres')

if errors:
    print(f'\n  IDs CIQUAL introuvables ({len(errors)}) :')
    for e in errors[:10]:
        print(f'    [{e["ig_id"]}] source_id={e["source_id"]}')

# Log JSON
log = DATA / 'nutrition/logs/patch_ciqual_axes.json'
log.parent.mkdir(parents=True, exist_ok=True)
log.write_text(json.dumps({
    'mode': mode, 'stats': dict(stats), 'by_axe': dict(by_axe),
    'patches': patches, 'errors': errors,
}, ensure_ascii=False, indent=2), encoding='utf-8')
print(f'\n  Log -> {log}')
if DRY_RUN:
    print('\n  DRY-RUN -- relancer avec --apply pour appliquer.')
