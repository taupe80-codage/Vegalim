"""
build_physical_v2.py
====================
Génère ingredient_physical.json en croisant ingredients_dictionary.json
avec les sources brutes de données nutritionnelles.

Sources (par priorité source dans dict_v2) :
  CIQUAL  → Table_Ciqual_2025_FR_2025_11_03.xlsx   (water_content_pct)
  USDA    → FoodData_Central_sr_legacy_food_csv_2018-04/  (water + portions)
             FoodData_Central_foundation_food_json_2025-12-18.json  (water + portions)
  CNF     → cnf/  (water_content_pct + edible_pct + portions)
  MANUAL  → pas de source brute → stub

Sortie :
  backend/data/ingredients/ingredient_physical.json

Usage :
  # Sans arguments : utilise les chemins par défaut dans backend/data/nutrition/raw/
  python scripts/nutrition/build_physical_v2.py

  # Avec chemins explicites :
  python scripts/nutrition/build_physical_v2.py \\
      --ciqual  <path>/Table_Ciqual_2025_FR_2025_11_03.xlsx \\
      --sr      <path>/FoodData_Central_sr_legacy_food_csv_2018-04/ \\
      --ff      <path>/FoodData_Central_foundation_food_json_2025-12-18.json \\
      --cnf     <path>/cnf/
"""

import json
import csv
import argparse
from datetime import datetime, timezone
from pathlib import Path

try:
    import openpyxl
except ImportError:
    raise SystemExit("openpyxl requis : pip install openpyxl")

UTC = timezone.utc

ROOT     = Path(__file__).parents[2]
DATA     = ROOT / 'backend/data'
RAW      = DATA / 'nutrition/raw'
DICT_V2  = DATA / 'ingredients/ingredients_dictionary.json'
PHYS_OUT = DATA / 'ingredients/ingredient_physical.json'

# Chemins par défaut des sources brutes (dans backend/data/nutrition/raw/)
_DEFAULT_CIQUAL = RAW / 'Table_Ciqual_2025_FR_2025_11_03.xlsx'
_DEFAULT_SR     = RAW / 'FoodData_Central_sr_legacy_food_csv_2018-04'
_DEFAULT_FF     = RAW / 'FoodData_Central_foundation_food_json_2025-12-18.json'
_DEFAULT_CNF    = RAW / 'cnf'

SCHEMA_VERSION = '2.0'


# ══════════════════════════════════════════════════════════════════
# INDEX CIQUAL
# ══════════════════════════════════════════════════════════════════

def build_ciqual_index(xlsx_path: Path) -> dict:
    """
    Retourne {str(alim_code): {'water_content_pct': float|None}}
    """
    wb = openpyxl.load_workbook(str(xlsx_path), data_only=True)
    ws = wb.active
    headers = [str(c.value).strip() if c.value else '' for c in next(ws.iter_rows(min_row=1, max_row=1))]
    alim_col = next(i for i, h in enumerate(headers) if 'alim_code' in h.lower())
    eau_col  = next((i for i, h in enumerate(headers) if 'eau' in h.lower()), None)

    idx = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        code = row[alim_col]
        if code is None:
            continue
        water = None
        if eau_col is not None:
            raw_w = row[eau_col]
            if raw_w not in (None, '-', '', 'traces'):
                try:
                    water = float(str(raw_w).replace(',', '.'))
                except ValueError:
                    pass
        idx[str(code)] = {'water_content_pct': water}
    return idx


# ══════════════════════════════════════════════════════════════════
# INDEX USDA SR LEGACY
# ══════════════════════════════════════════════════════════════════

def build_usda_sr_index(sr_dir: Path) -> dict:
    """
    Retourne {str(fdc_id): {'water_content_pct': float|None, 'portions': [...]}}
    """
    # Water = nutrient_id 1051
    water = {}
    with open(sr_dir / 'food_nutrient.csv', encoding='utf-8') as f:
        for row in csv.DictReader(f):
            if row.get('nutrient_id') == '1051':
                try:
                    water[row['fdc_id']] = float(row['amount'])
                except (KeyError, ValueError):
                    pass

    units_map = {}
    with open(sr_dir / 'measure_unit.csv', encoding='utf-8') as f:
        units_map = {row['id']: row['name'] for row in csv.DictReader(f)}

    portions: dict[str, list] = {}
    with open(sr_dir / 'food_portion.csv', encoding='utf-8') as f:
        for row in csv.DictReader(f):
            fid = row.get('fdc_id')
            gw  = row.get('gram_weight')
            amt = row.get('amount') or '1'
            desc = row.get('portion_description') or units_map.get(row.get('measure_unit_id', ''), '')
            if not fid or not gw:
                continue
            try:
                g_per_unit = float(gw) / (float(amt) or 1.0)
                if desc.strip():
                    portions.setdefault(fid, []).append({'label': desc.strip(), 'g': round(g_per_unit, 2)})
            except (ValueError, ZeroDivisionError):
                pass

    idx = {}
    with open(sr_dir / 'food.csv', encoding='utf-8') as f:
        for row in csv.DictReader(f):
            fid = row['fdc_id']
            idx[fid] = {
                'water_content_pct': water.get(fid),
                'portions':          portions.get(fid, []),
            }
    return idx


# ══════════════════════════════════════════════════════════════════
# INDEX USDA FOUNDATION FOODS
# ══════════════════════════════════════════════════════════════════

def build_usda_ff_index(ff_json: Path) -> dict:
    """
    Retourne {str(fdcId): {'water_content_pct': float|None, 'portions': [...]}}
    """
    with open(ff_json, encoding='utf-8') as f:
        ff_data = json.load(f)

    idx = {}
    for food in ff_data.get('FoundationFoods', []):
        fid = str(food.get('fdcId', ''))
        if not fid:
            continue
        water = None
        for n in food.get('foodNutrients', []):
            if n.get('nutrient', {}).get('name') == 'Water':
                try:
                    water = float(n['amount'])
                except (KeyError, ValueError):
                    pass
                break

        portions = []
        for p in food.get('foodPortions', []):
            label = p.get('measureUnit', {}).get('name', '') or p.get('modifier', '')
            gw    = p.get('gramWeight')
            amt   = p.get('amount') or p.get('value') or 1.0
            if label and gw:
                try:
                    portions.append({'label': label, 'g': round(float(gw) / float(amt), 2)})
                except (ValueError, ZeroDivisionError):
                    pass

        idx[fid] = {'water_content_pct': water, 'portions': portions}
    return idx


# ══════════════════════════════════════════════════════════════════
# INDEX CNF
# ══════════════════════════════════════════════════════════════════

def build_cnf_index(cnf_dir: Path) -> dict:
    """
    Retourne {str(FoodID): {'water_content_pct': float|None, 'edible_pct': float|None, 'portions': [...]}}
    """
    # Water = NutrientID 255
    water: dict[str, float] = {}
    with open(cnf_dir / 'NUTRIENT_AMOUNT.csv', encoding='latin-1') as f:
        reader = csv.DictReader(f)
        for row in reader:
            keys = list(row.keys())
            fid = row.get(keys[0])                    # FoodID (peut avoir BOM)
            nid = row.get('NutrientID') or (row.get(keys[1]) if len(keys) > 1 else None)
            if nid == '255' and fid:
                try:
                    water[str(fid)] = float(row.get('NutrientValue', 0))
                except ValueError:
                    pass

    # Edible % = 100 - RefuseAmount
    edible: dict[str, float] = {}
    with open(cnf_dir / 'REFUSE_AMOUNT.csv', encoding='latin-1') as f:
        for row in csv.DictReader(f):
            keys = list(row.keys())
            fid    = row.get(keys[0])
            refuse = row.get('RefuseAmount') or (row.get(keys[1]) if len(keys) > 1 else None)
            if fid and refuse:
                try:
                    edible[str(fid)] = round(100.0 - float(refuse), 2)
                except ValueError:
                    pass

    # Noms des mesures
    measure_names: dict[str, str] = {}
    with open(cnf_dir / 'MEASURE_NAME.csv', encoding='latin-1') as f:
        for row in csv.DictReader(f):
            keys = list(row.keys())
            mid  = row.get(keys[0])
            name = row.get('MeasureName') or (row.get(keys[1]) if len(keys) > 1 else '')
            if mid:
                measure_names[str(mid)] = name

    # Portions : ConversionFactorValue × 100 g
    portions: dict[str, list] = {}
    with open(cnf_dir / 'CONVERSION_FACTOR.csv', encoding='latin-1') as f:
        for row in csv.DictReader(f):
            keys   = list(row.keys())
            fid    = row.get(keys[0])
            mid    = row.get('MeasureID') or (row.get(keys[1]) if len(keys) > 1 else None)
            factor = row.get('ConversionFactorValue') or (row.get(keys[2]) if len(keys) > 2 else None)
            if fid and mid and factor:
                try:
                    g     = float(factor) * 100.0
                    label = measure_names.get(str(mid), str(mid))
                    portions.setdefault(str(fid), []).append({'label': label, 'g': round(g, 2)})
                except ValueError:
                    pass

    idx = {}
    with open(cnf_dir / 'FOOD_NAME.csv', encoding='latin-1') as f:
        for row in csv.DictReader(f):
            keys = list(row.keys())
            fid  = row.get(keys[0])
            if fid:
                idx[str(fid)] = {
                    'water_content_pct': water.get(str(fid)),
                    'edible_pct':        edible.get(str(fid)),
                    'portions':          portions.get(str(fid), []),
                }
    return idx


# ══════════════════════════════════════════════════════════════════
# HELPERS
# ══════════════════════════════════════════════════════════════════

def _default_units(entry: dict) -> tuple[str, dict]:
    axes = entry.get('axes', {})
    form = axes.get('form', '')
    if form in ('juice', 'liquid', 'oil', 'milk', 'cream', 'plant_water', 'syrup', 'vinegar'):
        return 'ml', {'ml': {'g': 1.0}, 'l': {'g': 1000.0}, 'g': {'g': 1}, 'kg': {'g': 1000}}
    return 'g', {'g': {'g': 1}, 'kg': {'g': 1000}}


_GENERIC_UNITS = {'g', 'kg', 'ml', 'l', 'gram', 'kilogram', ''}


def _inject_portions(units: dict, portions: list) -> None:
    for p in portions:
        label = p.get('label', '').strip().lower()
        if label and label not in _GENERIC_UNITS:
            units[label] = {'g': p['g']}


def _has_real_data(water, edible, portions) -> bool:
    return water is not None or edible is not None or bool(portions)


# ══════════════════════════════════════════════════════════════════
# COLLECT dict_v2 ENTRIES
# ══════════════════════════════════════════════════════════════════

def collect_dict_entries(categories: dict) -> dict:
    out = {}
    for _cat, cat in categories.items():
        for _sub, sub in cat.get('subcategories', {}).items():
            for key, entry in sub.get('ingredient_groups', {}).items():
                out[key] = entry
    return out


# ══════════════════════════════════════════════════════════════════
# MAIN
# ══════════════════════════════════════════════════════════════════

def main(args):
    print('=== BUILD PHYSICAL v2 ===\n')

    # ── Résolution des chemins (args ou défauts) ──────────────────
    ciqual_path = Path(args.ciqual) if args.ciqual else _DEFAULT_CIQUAL
    sr_path     = Path(args.sr)     if args.sr     else _DEFAULT_SR
    ff_path     = Path(args.ff)     if args.ff     else _DEFAULT_FF
    cnf_path    = Path(args.cnf)    if args.cnf    else _DEFAULT_CNF

    # ── Indexes ──────────────────────────────────────────────────
    print('Chargement des index sources...')
    ciqual_idx = build_ciqual_index(ciqual_path) if ciqual_path.exists() else {}
    print(f'  CIQUAL   : {len(ciqual_idx)} entries' + ('' if ciqual_path.exists() else ' ⚠ fichier absent'))

    sr_idx = build_usda_sr_index(sr_path) if sr_path.exists() else {}
    print(f'  USDA SR  : {len(sr_idx)} entries' + ('' if sr_path.exists() else ' ⚠ dossier absent'))

    ff_idx = build_usda_ff_index(ff_path) if ff_path.exists() else {}
    print(f'  USDA FF  : {len(ff_idx)} entries' + ('' if ff_path.exists() else ' ⚠ fichier absent'))

    # Merge SR + Foundation (Foundation plus récent en cas de collision)
    usda_idx = {**sr_idx, **ff_idx}
    print(f'  USDA tot : {len(usda_idx)} entries')

    cnf_idx = build_cnf_index(cnf_path) if cnf_path.exists() else {}
    print(f'  CNF      : {len(cnf_idx)} entries' + ('' if cnf_path.exists() else ' ⚠ dossier absent'))

    # ── dict_v2 ──────────────────────────────────────────────────
    print('\nChargement dict_v2...')
    dict_v2 = json.loads(DICT_V2.read_text(encoding='utf-8'))
    dict_entries = collect_dict_entries(dict_v2['categories'])
    print(f'  {len(dict_entries)} entrées\n')

    # ── Génération ───────────────────────────────────────────────
    stats = {'ciqual': 0, 'usda': 0, 'cnf': 0, 'manual': 0, 'stub': 0, 'ok': 0}
    output_entries = {}

    for dk, de in dict_entries.items():
        src = de.get('source', '')
        sid = str(de.get('source_id') or '')
        raw = None

        if src == 'CIQUAL' and sid in ciqual_idx:
            raw = ciqual_idx[sid]
            stats['ciqual'] += 1
        elif src == 'USDA' and sid in usda_idx:
            raw = usda_idx[sid]
            stats['usda'] += 1
        elif src == 'CNF' and sid in cnf_idx:
            raw = cnf_idx[sid]
            stats['cnf'] += 1
        elif src == 'MANUAL':
            stats['manual'] += 1
        else:
            stats['stub'] += 1

        default_unit, units = _default_units(de)
        units = dict(units)

        water   = raw.get('water_content_pct') if raw else None
        edible  = raw.get('edible_pct')        if raw else None
        portions = raw.get('portions', [])     if raw else []

        _inject_portions(units, portions)

        has_real = _has_real_data(water, edible, portions)
        if has_real:
            stats['ok'] += 1

        output_entries[dk] = {
            'default_unit':      default_unit,
            'units':             units,
            'density_g_per_ml':  None,
            'water_content_pct': water,
            'edible_pct':        edible,
            '_status':           'ok' if has_real else 'stub',
            '_tree_ing_id':      de.get('v32_ing_id'),
            '_tree_var_id':      de.get('v32_var_id'),
            '_dict_v2_key':      dk,
            '_source':           src,
            '_source_id':        de.get('source_id'),
            '_name_fr':          de.get('canonical_name_fr'),
            '_name_en':          de.get('canonical_name_en'),
            '_axes':             de.get('axes', {}),
        }

    total = len(output_entries)
    pct   = stats['ok'] / total * 100 if total else 0

    print(f'Résultats :')
    print(f'  CIQUAL   matched : {stats["ciqual"]}')
    print(f'  USDA     matched : {stats["usda"]}')
    print(f'  CNF      matched : {stats["cnf"]}')
    print(f'  MANUAL   stubs   : {stats["manual"]}')
    print(f'  Autres   stubs   : {stats["stub"]}')
    print(f'  → OK             : {stats["ok"]}/{total} ({pct:.1f}%)')

    # ── Écriture ─────────────────────────────────────────────────
    output = {
        '_meta': {
            'schema_version': SCHEMA_VERSION,
            'generated_at':   datetime.now(UTC).isoformat(),
            'generated_by':   'build_physical_v2.py',
            'built_from': [
                'ingredients_dictionary.json',
                'CIQUAL 2025',
                'USDA SR Legacy 2018',
                'USDA Foundation Foods 2025',
                'CNF 2015',
            ],
            'total_entries':  total,
            'entries_ok':     stats['ok'],
            'entries_stub':   total - stats['ok'],
            'match_stats': {
                'via_ciqual':     stats['ciqual'],
                'via_usda':       stats['usda'],
                'via_cnf':        stats['cnf'],
                'manual_stub':    stats['manual'],
                'unmatched_stub': stats['stub'],
            },
        }
    }
    output.update(output_entries)

    print(f'\nÉcriture → {PHYS_OUT}')
    tmp = PHYS_OUT.with_suffix('.tmp')
    tmp.write_text(json.dumps(output, ensure_ascii=False, indent=2), encoding='utf-8')
    tmp.replace(PHYS_OUT)

    size_kb = PHYS_OUT.stat().st_size // 1024
    print(f'\n{"="*50}')
    print(f'  BUILD PHYSICAL v2 — TERMINÉ')
    print(f'{"="*50}')
    print(f'  Total    : {total} entrées')
    print(f'  OK       : {stats["ok"]} ({pct:.1f}%)')
    print(f'  Stubs    : {total - stats["ok"]}')
    print(f'  Taille   : {size_kb} KB')


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description='Génère ingredient_physical.json depuis les sources brutes.')
    parser.add_argument('--ciqual', help='Chemin vers Table_Ciqual_2025_*.xlsx')
    parser.add_argument('--sr',     help='Dossier USDA SR Legacy CSV')
    parser.add_argument('--ff',     help='Chemin vers Foundation Foods JSON')
    parser.add_argument('--cnf',    help='Dossier CNF (avec NUTRIENT_AMOUNT.csv etc.)')
    main(parser.parse_args())
